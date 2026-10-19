#!/usr/bin/env python3
"""第 0–3 步的“表格底座”：初始化 / 校验 / 汇总全部资产表，输出 xlsx 总表 + 状态报告。可反复运行。

用法
  python scripts/build_tables.py --data data --init          # 新项目：建空表 + 拷入通用起步风格锁
  python scripts/build_tables.py --data data                 # 校验各表、补齐缺列、生成 out/资产总表.xlsx 和 out/状态报告.md
  python scripts/build_tables.py --config config/brand.yaml --data data   # 同时检查 brand.yaml 里还有哪些【待业务确认】

表的内容由 agent 按 steps/ 里的提示词填写（CSV，UTF-8-BOM，可直接导入飞书多维表/Excel）。本脚本不编造任何内容。
"""
import argparse, os, shutil, sys, datetime, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import common as C

ID_COL = {'爆款样本库_样本帖子': '样本ID', '爆款样本库_样本图片': '图片ID', '风格锁': '规则ID', '评审回流库': '回流ID',
          '发布队列': '队列ID', '投放映射': '计划名称', '自然流量库': '标题', '对标咨询证据': '匿名编号', '自有账号摘要': '账号'}


INBOX_README = '''收件箱：每篇帖子一个文件夹，放好后运行 python scripts/ingest_posts.py --data data
  bench/<任意名>/   对标帖     own/<任意名>/   自有帖
  文件夹里：01.jpg 02.jpg …（按图序）、caption.txt（第一行标题，后面正文）、
           metrics.txt（可选，「键: 值」：链接/作者/日期/点赞/收藏/评论/采集方式）、comments.txt（可选，「评论原文 | 意向理由」，不写用户名）
详见 steps/01-对标帖拆解.md
'''


def init(paths, cfg=None):
    for k in ('tables', 'exports', 'drafts', 'out', 'raw'):
        os.makedirs(paths[k], exist_ok=True)
    for k in ('bench', 'own'):
        os.makedirs(os.path.join(paths['root'], 'inbox', k), exist_ok=True)
    rd = os.path.join(paths['root'], 'inbox', 'README.txt')
    if not os.path.exists(rd): open(rd, 'w', encoding='utf-8').write(INBOX_README)
    made = []
    for name in C.TABLES:
        p = C.table_path(paths, name)
        if os.path.exists(p):
            continue
        if name == '风格锁':
            shutil.copy(os.path.join(C.TPL, '风格锁_通用起步.csv'), p)
            prof = C.load_profile(cfg or {}); off = set(prof.get('disable_rules') or [])
            if off:
                rows = C.load_table(paths, '风格锁')
                for r in rows:
                    if r['规则ID'] in off:
                        r['状态'] = f"停用·{prof['_id']}行业不适用"; r['专家意见'] = f"合规起步包（{prof['_id']}）自动停用：受监管行业不能用这类做法"
                C.save_table(paths, '风格锁', rows)
            n_on = sum(1 for r in C.load_table(paths, '风格锁') if r.get('状态', '').startswith('生效'))
            print(f'风格锁：{n_on} 条通用起步规则已自动生效' + (f"，{len(off)} 条因{prof['_id']}行业停用（{'、'.join(sorted(off))}）" if off else '')
                  + '。一句话：像小红书笔记、不像广告；事实只用已确认的；价格不当钩子；主体上的字要清楚。不用逐条确认，评审时遇到再调。')
        else:
            C.save_table(paths, name, [])
        made.append(name)
    print('已初始化：', paths['root'], '新建表：', '、'.join(made) or '无（都已存在）')


def unknowns(obj, prefix=''):
    out = []
    if isinstance(obj, dict):
        for k, v in obj.items():
            out += unknowns(v, f'{prefix}.{k}' if prefix else str(k))
    elif isinstance(obj, list):
        for i, v in enumerate(obj):
            out += unknowns(v, f'{prefix}.{i}')
    elif C.is_unknown(obj) and obj is not None and str(obj).strip():
        out.append(prefix)
    return out


def check(paths, cfg=None, where=None):
    rep, warn, data = [], [], {}
    for name in C.TABLES:
        p = C.table_path(paths, name)
        if not os.path.exists(p):
            warn.append(f'缺表 {name}（先运行 --init）'); continue
        rows = C.read_csv(p); cols = C.schema(name)
        have = list(rows[0].keys()) if rows else cols
        extra = [c for c in have if c not in cols]
        if any(c not in have for c in cols) or not rows:
            # 补齐缺列，保留额外列
            import csv
            with open(p, encoding='utf-8-sig') as f:
                hdr = next(csv.reader(f), [])
            missing = [c for c in cols if c not in hdr]
            if missing:
                C.write_csv(p, rows, cols + extra); warn.append(f'{name}：补齐缺列 {missing}')
        if extra:
            warn.append(f'{name}：有模板外的列 {extra}（保留，不影响脚本）')
        idc = ID_COL.get(name)
        if idc:
            dup = [k for k, n in collections.Counter(r.get(idc, '') for r in rows).items() if k and n > 1]
            if dup: warn.append(f'{name}：{idc} 重复 {dup}')
            blank = sum(1 for r in rows if not r.get(idc, '').strip())
            if blank: warn.append(f'{name}：{blank} 行没有 {idc}')
        data[name] = rows
    # ---- 业务规则检查（只提示，不改内容）
    sl = data.get('风格锁', [])
    eff = [r for r in sl if r.get('状态', '').startswith('生效')]
    prop = [r for r in sl if r.get('状态', '').startswith('提议')]
    rv = data.get('评审回流库', [])
    pending = [r for r in rv if r.get('是否已回写', '') not in ('是', '不需要')]
    q = data.get('发布队列', [])
    qs = collections.Counter(r.get('状态', '') for r in q)
    for r in q:
        if r.get('状态') in ('已通过，待发布', '已发布') and r.get('专家结论', '') not in ('通过',):
            warn.append(f"发布队列 {r.get('队列ID')}：状态是「{r.get('状态')}」但专家结论不是「通过」——人工关卡没过不能发")
        if r.get('含AI内容') == '是' and '含AI' not in (r.get('待确认事项', '') + r.get('备注', '')):
            warn.append(f"发布队列 {r.get('队列ID')}：含AI内容=是，发布时记得在小红书勾选「含AI生成内容」")
    bench = [r for r in data.get('爆款样本库_样本帖子', []) if r.get('来源类型') == '对标']
    own = [r for r in data.get('爆款样本库_样本帖子', []) if r.get('来源类型', '').startswith('自有')]
    imgs = data.get('爆款样本库_样本图片', [])
    unk_cells = sum(1 for n, rows in data.items() if n != '风格锁' for r in rows for v in r.values() if v and C.UNKNOWN in str(v))
    todo_cells = sum(1 for n in ('爆款样本库_样本帖子', '爆款样本库_样本图片') for r in data.get(n, []) for v in r.values() if v == C.TODO)
    rep += [f'# 状态报告（{datetime.datetime.now():%Y-%m-%d %H:%M}）', '']
    if where: rep += [f'**{where[0]}**', '', f'- {where[1]}', f'- {where[2]}', '']
    rep += [f'数据目录：`{paths["root"]}`', '',
            '| 表 | 行数 |', '|---|---|'] + [f'| {n} | {len(data.get(n, []))} |' for n in C.TABLES]
    rep += ['', f'- 样本：对标 {len(bench)} 篇，自有 {len(own)} 篇，图片拆解 {len(imgs)} 张',
            f'- 风格锁：生效 {len(eff)} 条，提议（未生效、待专家确认）{len(prop)} 条',
            f'- 评审回流：共 {len(rv)} 条，待回写 {len(pending)} 条' + (f"（{'、'.join(r.get('回流ID','') for r in pending)}）" if pending else ''),
            '- 发布队列：' + ('，'.join(f'{k or "空"} {v}' for k, v in qs.items()) or '空'),
            f'- 表里还有 {unk_cells} 处【待业务确认】',
            f'- 样本库还有 {todo_cells} 格{C.TODO}（ingest_posts.py 建的空壳，按 templates/拆解模板.md 填）',
            f'- 对标咨询证据 {len(data.get("对标咨询证据", []))} 条；自有账号摘要 {len(data.get("自有账号摘要", []))} 个账号']
    if cfg is not None:
        import facts as F
        u = [k for k in unknowns(cfg) if not F.irrelevant(cfg, k)]   # 没开聚光时 juguang.* 之类不算待确认
        rep.append(f'- brand.yaml 还有 {len(u)} 项【待业务确认】' + (f'：{", ".join(u[:20])}' + ('…' if len(u) > 20 else '') if u else ''))
    if warn:
        rep += ['', '## 提示'] + [f'- {w}' for w in warn]
    return data, rep


def to_xlsx(data, path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); wb.remove(wb.active)
    for name in C.TABLES:
        rows = data.get(name)
        if rows is None: continue
        ws = wb.create_sheet(name[:31]); cols = list(rows[0].keys()) if rows else C.schema(name)
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='3A3A3A')
        for r in rows: ws.append([r.get(c, '') for c in cols])
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(40, max(10, len(c) * 2 + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row: c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'B2'
    os.makedirs(os.path.dirname(path), exist_ok=True); wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--data', default='data'); ap.add_argument('--config'); ap.add_argument('--init', action='store_true')
    a = ap.parse_args(); paths = C.data_paths(a.data)
    cfg = C.load_yaml(a.config) if a.config else None
    if a.init: init(paths, cfg)
    where = None
    if a.config:
        try:
            import xhs
            st = xhs.state(argparse.Namespace(config=a.config, data=a.data))
            n, pending, nxt = st[0], st[1], st[2]; ex = st[5] if len(st) > 5 else ''
            where = (f'现在：{xhs.tag(n, ex)}', f'还差：{pending}', f'下一步：{nxt}')
        except SystemExit:
            pass
    data, rep = check(paths, cfg, where)
    to_xlsx(data, os.path.join(paths['out'], '资产总表.xlsx'))
    open(os.path.join(paths['out'], '状态报告.md'), 'w', encoding='utf-8').write('\n'.join(rep) + '\n')
    print('\n'.join(rep)); print(f"\n输出：{paths['out']}/资产总表.xlsx、状态报告.md")


if __name__ == '__main__':
    main()
